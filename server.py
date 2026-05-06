import http.server
from http.server import HTTPServer, BaseHTTPRequestHandler
import json
import os
import csv
import traceback
from datetime import datetime, date
from urllib.parse import urlparse, parse_qs
from collections import defaultdict
import threading
import time
import base64
import re

import config
from automail import load_contacts, generate_email, send_email, log_to_csv, get_sent_emails, check_bounces, validate_email, log, debug, _file_logger

# ═══════════════════════════════════════════════════════════════════
# SERVER HANDLER
# ═══════════════════════════════════════════════════════════════════

class AutomailHandler(BaseHTTPRequestHandler):
    is_sending = False
    is_paused = False
    stop_requested = False
    send_progress = {"current": 0, "total": 0, "status": "idle"}

    def send_json(self, data, status=200):
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(json.dumps(data).encode("utf-8"))

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/" or parsed.path == "/dashboard.html":
            self.serve_file("dashboard.html", "text/html")
        elif parsed.path == "/api/contacts":
            self.handle_get_contacts(parsed)
        elif parsed.path == "/api/stats":
            self.handle_get_stats()
        elif parsed.path == "/api/resumes":
            self.handle_get_resumes()
        elif parsed.path == "/api/logs":
            self.handle_get_logs()
        elif parsed.path == "/api/progress":
            self.send_json({**AutomailHandler.send_progress, "paused": AutomailHandler.is_paused})
        elif parsed.path == "/api/settings":
            self.send_json({
                "gmail_address": config.GMAIL_ADDRESS,
                "gmail_app_password": config.GMAIL_APP_PASSWORD,
                "delay": config.DELAY_BETWEEN_EMAILS,
                "batch_size": config.BATCH_SIZE,
                "target_role": config.TARGET_ROLE
            })
        elif parsed.path == "/api/export-logs":
            self.handle_export_logs()
        elif parsed.path == "/api/health":
            self.handle_health_check()
        elif parsed.path == "/api/debug-log":
            self.handle_debug_log(parsed)
        elif parsed.path == "/api/templates":
            self.handle_get_templates()
        elif parsed.path == "/api/preview":
            self.handle_preview(parsed)
        else:
            super().do_GET()

    def do_POST(self):
        content_length = int(self.headers.get("Content-Length", 0))
        body = self.rfile.read(content_length).decode("utf-8") if content_length else "{}"
        try: data = json.loads(body) if body else {}
        except: data = {}

        parsed = urlparse(self.path)
        if parsed.path == "/api/send-batch":
            self.handle_send_batch(data)
        elif parsed.path == "/api/retry-failed":
            self.handle_retry_failed(data)
        elif parsed.path == "/api/pause":
            AutomailHandler.is_paused = not AutomailHandler.is_paused
            self.send_json({"success": True, "paused": AutomailHandler.is_paused})
        elif parsed.path == "/api/stop":
            AutomailHandler.stop_requested = True
            self.send_json({"success": True})
        elif parsed.path == "/api/clear-logs":
            if os.path.exists(config.SEND_LOG_PATH):
                try:
                    with open(config.SEND_LOG_PATH, "w", encoding="utf-8", newline="") as f:
                        writer = csv.writer(f)
                        writer.writerow(["Timestamp", "Sr.No", "Company", "Email", "Name", "Subject", "Resume", "Status"])
                except: pass
            self.send_json({"success": True})
        elif parsed.path == "/api/settings":
            self.handle_settings_post(data)
        elif parsed.path == "/api/upload-contacts":
            self.handle_upload_contacts(data)
        elif parsed.path == "/api/add-contact":
            self.handle_add_contact(data)
        elif parsed.path == "/api/edit-contact":
            self.handle_edit_contact(data)
        elif parsed.path == "/api/delete-contact":
            self.handle_delete_contact(data)
        elif parsed.path == "/api/templates":
            self.handle_save_templates(data)
        elif parsed.path == "/api/upload-template-asset":
            self.handle_upload_template_asset(data)
        else:
            self.send_json({"error": "Not found"}, 404)

    def handle_upload_template_asset(self, data):
        filename = data.get("filename")
        content_b64 = data.get("content")
        if not filename or not content_b64:
            return self.send_json({"error": "Missing data"}, 400)
        
        filename = re.sub(r'[^a-zA-Z0-9._-]', '_', filename)
        asset_dir = os.path.join(config.BASE_DIR, "template_assets")
        if not os.path.exists(asset_dir):
            os.makedirs(asset_dir)
        path = os.path.join(asset_dir, filename)
        
        try:
            with open(path, "wb") as f:
                f.write(base64.b64decode(content_b64))
            self.send_json({"success": True, "path": f"template_assets/{filename}"})
        except Exception as e:
            self.send_json({"error": str(e)}, 500)

    def handle_get_templates(self):
        from automail import load_templates
        self.send_json({"templates": load_templates()})

    def handle_save_templates(self, data):
        from automail import save_templates
        save_templates(data.get("templates", []))
        self.send_json({"success": True})

    def handle_preview(self, parsed):
        from automail import load_contacts, load_templates, render_template
        params = parse_qs(parsed.query)
        idx = int(params.get("idx", [0])[0])
        sr = int(params.get("sr", [1])[0])
        role = params.get("role", [config.TARGET_ROLE])[0]
        
        contacts = load_contacts()
        contact = next((c for c in contacts if c["sr"] == sr), contacts[0] if contacts else {})
        templates = load_templates()
        template = templates[idx] if idx < len(templates) else templates[0]
        
        rendered = render_template(template, contact, role)
        self.send_json(rendered)

    def serve_file(self, filename, content_type):
        try:
            with open(os.path.join(config.BASE_DIR, filename), "rb") as f:
                self.send_response(200)
                self.send_header("Content-type", content_type)
                self.end_headers()
                self.wfile.write(f.read())
        except Exception:
            self.send_error(404, "File not found")

    def handle_get_contacts(self, parsed):
        params = parse_qs(parsed.query)
        start = int(params.get("start", [1])[0])
        count = int(params.get("count", [50])[0])
        search = params.get("search", [""])[0].lower()
        
        contacts = load_contacts(start=1, count=99999)
        if search:
            contacts = [c for c in contacts if search in c["company"].lower() or search in c["email"].lower() or search in c["name"].lower()]
        
        total_filtered = len(contacts)
        contacts = contacts[start-1:start-1+count]
        self.send_json({"contacts": contacts, "total": total_filtered})

    def handle_get_stats(self):
        contacts = load_contacts(start=1, count=99999)
        sent_count = 0
        failed_count = 0
        bounced_count = 0
        today_sent = 0
        today_str = date.today().isoformat()
        
        history_by_day = defaultdict(int)
        
        if os.path.exists(config.SEND_LOG_PATH):
            with open(config.SEND_LOG_PATH, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    status = row.get("Status", "")
                    ts = row.get("Timestamp", "")
                    day = ts[:10]
                    if "SENT" in status:
                        sent_count += 1
                        history_by_day[day] += 1
                        if day == today_str: today_sent += 1
                    elif "FAILED" in status:
                        failed_count += 1
                        if "Bounced" in status: bounced_count += 1
        
        sorted_days = sorted(history_by_day.keys())[-7:]
        daily_stats = [{"day": d, "count": history_by_day[d]} for d in sorted_days]

        self.send_json({
            "total_leads": len(contacts),
            "sent_count": sent_count,
            "failed_count": failed_count,
            "bounced_count": bounced_count,
            "today_sent": today_sent,
            "daily_limit": config.MAX_DAILY_EMAILS,
            "daily_stats": daily_stats
        })

    def handle_get_resumes(self):
        resumes = []
        for key, info in config.AVAILABLE_RESUMES.items():
            exists = os.path.exists(info["path"])
            resumes.append({"key": key, "label": info["label"], "exists": exists})
        self.send_json({"resumes": resumes})

    def handle_get_logs(self):
        logs = []
        if os.path.exists(config.SEND_LOG_PATH):
            try:
                with open(config.SEND_LOG_PATH, "r", encoding="utf-8") as f:
                    logs = list(csv.DictReader(f))
            except: pass
        self.send_json({"logs": logs[::-1][:200]})

    def handle_export_logs(self):
        if not os.path.exists(config.SEND_LOG_PATH):
            self.send_error(404, "No logs found")
            return
        with open(config.SEND_LOG_PATH, "rb") as f:
            self.send_response(200)
            self.send_header("Content-Type", "text/csv")
            self.send_header("Content-Disposition", 'attachment; filename="automail_history.csv"')
            self.end_headers()
            self.wfile.write(f.read())

    def handle_settings_post(self, data):
        try:
            config.GMAIL_ADDRESS = data.get("gmail_address", config.GMAIL_ADDRESS)
            config.GMAIL_APP_PASSWORD = data.get("gmail_app_password", config.GMAIL_APP_PASSWORD)
            config.DELAY_BETWEEN_EMAILS = int(data.get("delay", config.DELAY_BETWEEN_EMAILS))
            config.TARGET_ROLE = data.get("target_role", config.TARGET_ROLE)
            
            env_path = os.path.join(config.BASE_DIR, ".env")
            with open(env_path, "w") as f:
                f.write(f"GMAIL_ADDRESS={config.GMAIL_ADDRESS}\n")
                f.write(f"GMAIL_APP_PASSWORD={config.GMAIL_APP_PASSWORD}\n")
                f.write(f"TARGET_ROLE=\"{config.TARGET_ROLE}\"\n")
            self.send_json({"success": True})
        except Exception as e:
            self.send_json({"error": str(e)}, 500)

    def handle_upload_contacts(self, data):
        filename = data.get("filename", "contacts.csv")
        b64data = data.get("data", "")
        if not b64data: return self.send_json({"error": "No data"}, 400)
        try:
            if "," in b64data: b64data = b64data.split(",")[-1]
            filepath = os.path.join(config.BASE_DIR, filename)
            with open(filepath, "wb") as f:
                f.write(base64.b64decode(b64data))
            config.HR_EXCEL_PATH = filepath
            cfg_path = os.path.join(config.BASE_DIR, "config.py")
            with open(cfg_path, "r") as f: content = f.read()
            content = re.sub(r'HR_EXCEL_PATH\s*=\s*.*', f'HR_EXCEL_PATH = os.path.join(BASE_DIR, "{filename}")', content)
            with open(cfg_path, "w") as f: f.write(content)
            self.send_json({"success": True})
        except Exception as e:
            self.send_json({"error": str(e)}, 500)

    def handle_add_contact(self, data):
        name, company, email = data.get("name"), data.get("company"), data.get("email")
        if not all([name, company, email]): return self.send_json({"error": "Missing fields"}, 400)
        filepath = config.HR_EXCEL_PATH
        try:
            if filepath.lower().endswith(".csv"):
                with open(filepath, "a", newline="", encoding="utf-8") as f:
                    writer = csv.reader(open(filepath, "r", encoding="utf-8"))
                    rows = list(writer)
                    next_sr = len(rows) + 1 if rows else 1
                    csv.writer(f).writerow([next_sr, company, email, name])
            else:
                import openpyxl
                wb = openpyxl.load_workbook(filepath)
                ws = wb.active
                ws.append([ws.max_row, company, email, name])
                wb.save(filepath)
            self.send_json({"success": True})
        except Exception as e:
            self.send_json({"error": str(e)}, 500)

    def handle_edit_contact(self, data):
        sr, name, company, email = data.get("sr"), data.get("name"), data.get("company"), data.get("email")
        filepath = config.HR_EXCEL_PATH
        try:
            if filepath.lower().endswith(".csv"):
                rows = []
                with open(filepath, "r", encoding="utf-8") as f: rows = list(csv.reader(f))
                for r in rows:
                    if r and str(r[0]) == str(sr):
                        if len(r) > 1: r[1] = company
                        if len(r) > 2: r[2] = email
                        if len(r) > 3: r[3] = name
                with open(filepath, "w", newline="", encoding="utf-8") as f: csv.writer(f).writerows(rows)
            else:
                import openpyxl
                wb = openpyxl.load_workbook(filepath)
                ws = wb.active
                for row in ws.iter_rows():
                    if str(row[0].value) == str(sr):
                        row[1].value, row[2].value, row[3].value = company, email, name
                wb.save(filepath)
            self.send_json({"success": True})
        except Exception as e:
            self.send_json({"error": str(e)}, 500)

    def handle_delete_contact(self, data):
        sr = data.get("sr")
        filepath = config.HR_EXCEL_PATH
        try:
            if filepath.lower().endswith(".csv"):
                rows = []
                with open(filepath, "r", encoding="utf-8") as f: rows = list(csv.reader(f))
                new_rows = [r for r in rows if r and str(r[0]) != str(sr)]
                with open(filepath, "w", newline="", encoding="utf-8") as f: csv.writer(f).writerows(new_rows)
            else:
                import openpyxl
                wb = openpyxl.load_workbook(filepath)
                ws = wb.active
                for idx, row in enumerate(ws.iter_rows(), 1):
                    if str(row[0].value) == str(sr):
                        ws.delete_rows(idx); break
                wb.save(filepath)
            self.send_json({"success": True})
        except Exception as e:
            self.send_json({"error": str(e)}, 500)

    def handle_send_batch(self, data):
        if AutomailHandler.is_sending: return self.send_json({"error": "Batch already running"}, 400)
        start = int(data.get("start", 1))
        count = int(data.get("count", 50))
        role = data.get("role", config.TARGET_ROLE)
        resume = data.get("resume", "vinoth_main")
        threading.Thread(target=self._batch_worker, args=(start, count, role, resume), daemon=True).start()
        self.send_json({"success": True})

    def handle_retry_failed(self, data):
        if AutomailHandler.is_sending: return self.send_json({"error": "Busy"}, 400)
        role = data.get("role", config.TARGET_ROLE)
        resume = data.get("resume", "vinoth_main")
        threading.Thread(target=self._batch_worker, args=(1, 9999, role, resume, True), daemon=True).start()
        self.send_json({"success": True})

    def handle_health_check(self):
        import smtplib
        result = {"smtp": False, "smtp_error": "", "imap": False, "imap_error": ""}
        if config.GMAIL_ADDRESS and config.GMAIL_APP_PASSWORD:
            try:
                with smtplib.SMTP(config.SMTP_SERVER, config.SMTP_PORT, timeout=8) as server:
                    server.ehlo(); server.starttls()
                    server.login(config.GMAIL_ADDRESS, config.GMAIL_APP_PASSWORD)
                result["smtp"] = True
            except Exception as e: result["smtp_error"] = str(e)
            try:
                import imaplib
                mail = imaplib.IMAP4_SSL("imap.gmail.com")
                mail.login(config.GMAIL_ADDRESS, config.GMAIL_APP_PASSWORD)
                mail.logout()
                result["imap"] = True
            except Exception as e: result["imap_error"] = str(e)
        self.send_json(result)

    def handle_debug_log(self, parsed):
        params = parse_qs(parsed.query)
        lines_n = int(params.get("lines", [100])[0])
        log_path = os.path.join(config.BASE_DIR, "debug.log")
        if not os.path.exists(log_path):
            self.send_json({"lines": []}); return
        try:
            with open(log_path, "r", encoding="utf-8", errors="replace") as f:
                all_lines = f.readlines()
            tail = [l.rstrip() for l in all_lines[-lines_n:]]
            self.send_json({"lines": tail})
        except: self.send_json({"lines": ["Error reading log"]})

    def _batch_worker(self, start, count, role, resume_key, retry_failed=False):
        AutomailHandler.is_sending = True
        AutomailHandler.is_paused = False
        AutomailHandler.stop_requested = False
        try:
            all_contacts = load_contacts(start=1, count=99999)
            if retry_failed:
                failed_emails = set()
                if os.path.exists(config.SEND_LOG_PATH):
                    with open(config.SEND_LOG_PATH, "r") as f:
                        for r in csv.DictReader(f):
                            if "FAILED" in r.get("Status", ""): failed_emails.add(r.get("Email", ""))
                contacts = [c for c in all_contacts if c["email"] in failed_emails]
            else:
                sent = get_sent_emails()
                contacts = all_contacts[start-1 : start-1+count]
                contacts = [c for c in contacts if c["email"] not in sent]
            
            AutomailHandler.send_progress = {"current": 0, "total": len(contacts), "status": "running"}
            resume_path = config.AVAILABLE_RESUMES.get(resume_key, {}).get("path", config.DEFAULT_RESUME_PATH)
            
            for i, c in enumerate(contacts):
                while AutomailHandler.is_paused and not AutomailHandler.stop_requested: time.sleep(1)
                if AutomailHandler.stop_requested: break
                AutomailHandler.send_progress["current"] = i + 1
                AutomailHandler.send_progress["status"] = f"Sending to {c['company']}..."
                
                email_data = generate_email(c, role)
                current_attachments = email_data.get("attachments", [])
                if not current_attachments: current_attachments = [resume_path]
                else:
                    current_attachments = [(os.path.join(config.BASE_DIR, p) if not os.path.isabs(p) else p) for p in current_attachments]

                res = send_email(c["email"], email_data["subject"], email_data["body"], current_attachments)
                
                if res is True: status = "SENT"
                else:
                    status = f"FAILED ({str(res)})"
                    AutomailHandler.send_progress["last_error"] = str(res)
                
                log_to_csv(c, email_data["subject"], status, os.path.basename(resume_path))
                if i < len(contacts) - 1: time.sleep(config.DELAY_BETWEEN_EMAILS)
            
            AutomailHandler.send_progress["status"] = "completed"
        except Exception as e:
            AutomailHandler.send_progress["status"] = f"error: {str(e)}"
            traceback.print_exc()
        finally: AutomailHandler.is_sending = False

def main():
    def bounce_loop():
        while True:
            try: check_bounces()
            except: pass
            time.sleep(300)
    threading.Thread(target=bounce_loop, daemon=True).start()
    port = 8000
    server = HTTPServer(("localhost", port), AutomailHandler)
    print(f"\n⚡ AUTOMAIL Dashboard: http://localhost:{port}\n")
    try: server.serve_forever()
    except KeyboardInterrupt: server.server_close()

if __name__ == "__main__": main()
